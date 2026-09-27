"""일일 요약에 첨부하는 엑셀.

메일 본문은 오늘 볼 것만 담고(신규·마감임박), 전체 목록과 수집 현황은 이 파일로
보낸다. 본문에 다 넣으면 길어서 정작 급한 건을 놓친다.

D-day 는 값이 아니라 수식으로 넣는다. 받은 편지함에 며칠 묵혔다 열어도 그날
기준으로 다시 세도록 — 값으로 박으면 열었을 때 이미 틀린 숫자다.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import ConfigBundle
from ..extract.deadline import KST
from ..models import FIELD_NAMES, Posting
from ..store import Store
from .build import ReportData, deadline_str  # noqa: F401  (표시 형식을 맞추려고 같은 모듈을 쓴다)

FONT = "Arial"
INK = "191F28"
HEAD_BG = "1B64DA"
URGENT_BG = "FFF0F1"
SOON_BG = "FFF7E6"
THIN = Side(style="thin", color="D7DCE3")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

COLS = ["No", "분야", "기관명", "지역", "제목", "게시일", "마감일", "D-day", "점수", "출처"]
WIDTHS = [5, 12, 24, 16, 60, 12, 12, 9, 7, 26]


def _head(ws, row: int, cols: list[str]) -> None:
    for i, name in enumerate(cols, start=1):
        c = ws.cell(row, i, name)
        c.font = Font(name=FONT, bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD_BG)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24
    ws.freeze_panes = ws.cell(row + 1, 1)


def _widths(ws, spec: list[int]) -> None:
    for i, w in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _posting_sheet(wb: Workbook, title: str, rows: list[Posting], names: dict[str, str],
                   now: datetime) -> None:
    ws = wb.create_sheet(f"{title} ({len(rows)})"[:31])
    ws.cell(1, 1, f"{title} · {len(rows)}건").font = Font(name=FONT, bold=True, size=13, color=INK)
    ws.cell(2, 1, f"기준 {now.astimezone(KST):%Y-%m-%d %H:%M} KST · D-day 는 파일을 여는 날짜로 다시 계산됩니다").font = Font(
        name=FONT, size=9, color="8B95A1")
    _head(ws, 4, COLS)

    r = 5
    for i, p in enumerate(rows, start=1):
        src = p.sources[0] if p.sources else None
        ws.cell(r, 1, i)
        ws.cell(r, 2, FIELD_NAMES.get(p.field, p.field))
        ws.cell(r, 3, p.org_name)
        ws.cell(r, 4, ", ".join(p.region) if p.region else "")
        cell = ws.cell(r, 5, p.title)
        if src:
            cell.hyperlink = src.url
            cell.font = Font(name=FONT, size=10, color="1B64DA", underline="single")
        if p.posted_at:
            # posted_at 은 date, deadline 은 datetime 이다. 같이 다루면 터진다
            ws.cell(r, 6, p.posted_at).number_format = "yyyy-mm-dd"
        if p.deadline:
            ws.cell(r, 7, p.deadline.astimezone(KST).replace(tzinfo=None)).number_format = "yyyy-mm-dd"
            ws.cell(r, 8, f'=IF(G{r}="","",INT(G{r})-TODAY())').number_format = '"D-"0;"지남 "0"일";"D-DAY"'
        else:
            ws.cell(r, 7, "원문확인")
            ws.cell(r, 8, "원문확인")
        ws.cell(r, 9, p.relevance_score)
        ws.cell(r, 10, names.get(src.source_id, src.source_id) if src else "")

        left = (p.deadline.astimezone(KST).date() - now.astimezone(KST).date()).days if p.deadline else None
        for c in range(1, len(COLS) + 1):
            cell = ws.cell(r, c)
            if c != 5 or not src:
                cell.font = Font(name=FONT, size=10, color=INK)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 5),
                                       horizontal="center" if c in (1, 6, 7, 8, 9) else "left")
            if left is not None and 0 <= left <= 1:
                cell.fill = PatternFill("solid", fgColor=URGENT_BG)
            elif left is not None and left <= 3:
                cell.fill = PatternFill("solid", fgColor=SOON_BG)
        ws.row_dimensions[r].height = 28
        r += 1

    if rows:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(COLS))}{r - 1}"
    else:
        ws.cell(5, 1, "해당 없음").font = Font(name=FONT, size=10, color="8B95A1")
    _widths(ws, WIDTHS)


def _status_sheet(wb: Workbook, data: ReportData, bundle: ConfigBundle, now: datetime) -> None:
    ws = wb.create_sheet("수집현황")
    ws.cell(1, 1, "수집현황").font = Font(name=FONT, bold=True, size=13, color=INK)
    ws.cell(2, 1, "'중단' 은 기술 문제가 아니라 상대 사이트의 정책·차단 때문입니다").font = Font(
        name=FONT, size=9, color="8B95A1")

    cols = ["소스", "지역", "상태", "목록 글 수", "신규", "비고", "주소"]
    _head(ws, 4, cols)

    by_id = {s.source_id: s for s in (data.run.sources if data.run else [])}
    STATUS = {"ok": "정상", "empty": "0건", "fail": "실패", "unconfigured": "미설정"}

    r = 5
    for s in sorted(bundle.sources, key=lambda x: (not x.enabled, x.id)):
        res = by_id.get(s.id)
        ws.cell(r, 1, s.name)
        ws.cell(r, 2, ", ".join(s.region_hint or []))
        if not s.enabled:
            ws.cell(r, 3, "중단")
            ws.cell(r, 6, (s.notes or "").split(".")[0][:120])
        else:
            ws.cell(r, 3, STATUS.get(res.status, res.status) if res else "미실행")
            if res:
                ws.cell(r, 4, res.listed)
                ws.cell(r, 5, res.new)
                if res.errors:
                    ws.cell(r, 6, res.errors[0][:120])
        ws.cell(r, 7, s.adapter.get("list_url") or s.adapter.get("endpoint") or "")
        for c in range(1, len(cols) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT, size=10, color=INK)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (1, 6, 7)),
                                       horizontal="center" if c in (3, 4, 5) else "left")
            if not s.enabled:
                cell.fill = PatternFill("solid", fgColor="F2F4F6")
            elif res and res.status == "fail":
                cell.fill = PatternFill("solid", fgColor=URGENT_BG)
        ws.row_dimensions[r].height = 30
        r += 1

    ws.auto_filter.ref = f"A4:G{r - 1}"
    _widths(ws, [40, 16, 10, 11, 8, 46, 52])


def build_workbook(data: ReportData, store: Store, bundle: ConfigBundle, now: datetime) -> bytes:
    """메일에 붙일 엑셀을 바이트로 돌려준다."""
    wb = Workbook()
    wb.remove(wb.active)

    active = [p for p in store.values()
              if p.status.value != "expired"
              and not (p.deadline and p.deadline < now)
              and "피드백제외" not in p.flags]
    active.sort(key=lambda p: (p.deadline.timestamp() if p.deadline else 9e18, -(p.relevance_score or 0)))

    _posting_sheet(wb, "신규", data.new, data.source_names, now)
    _posting_sheet(wb, "마감임박", data.closing, data.source_names, now)
    _posting_sheet(wb, "진행중 전체", active, data.source_names, now)
    _status_sheet(wb, data, bundle, now)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def workbook_name(now: datetime) -> str:
    return f"부울경_평생교육원_강사공고_{now.astimezone(KST):%Y-%m-%d}.xlsx"
