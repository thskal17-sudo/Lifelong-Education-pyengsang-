"""강사잇다 업로드 양식으로 공고를 담는다.

사이트가 받는 양식이라 열 이름·색·툴팁이 정확해야 한다. 우리가 표를 흉내 내는
대신 받은 템플릿 파일(assets/gangsaitda_template.xlsx)을 열어 '공고' 시트만
채운다 — 그래야 '예시'·'안내' 시트와 머리글 서식이 그대로 간다.

담는 기준은 양식 안내를 따른다: 마감일이 지난 줄은 올라가지 않으므로 넣지 않고,
한 줄에 공고 하나씩 적는다.
"""
from __future__ import annotations

from datetime import datetime
from importlib.resources import files
from io import BytesIO

from openpyxl import load_workbook

from ..extract.deadline import KST
from ..models import Posting
from ..store import Store

SHEET = "공고"
# 양식의 열 순서. 이름을 고치면 사이트가 못 읽으므로 템플릿에서 읽어 대조한다
COLUMNS = [
    "제목", "기관명", "지역", "마감일", "수업 일정", "상세 내용",
    "수업 대상", "모집 인원", "지원 자격", "제출 서류",
    "원문 링크", "지원서 링크", "지원 이메일",
]


def _template_bytes() -> bytes:
    return (files("gia.report.assets") / "gangsaitda_template.xlsx").read_bytes()


def uploadable(store: Store, now: datetime) -> list[Posting]:
    """올릴 수 있는 공고. 마감이 지났거나 없는 것은 뺀다.

    양식 안내에 '이미 지난 날짜는 올라가지 않아요' 라고 적혀 있다. 마감일이 없는
    공고도 뺀다 — 마감일은 파란 칸(필수)이라 비우면 그 줄이 통째로 거부된다.
    """
    out = [p for p in store.values()
           if p.status.value != "expired"
           and p.deadline is not None
           and p.deadline >= now
           and "피드백제외" not in p.flags]
    out.sort(key=lambda p: (p.deadline, -(p.relevance_score or 0)))
    return out


def _row(p: Posting) -> list:
    url = p.sources[0].url if p.sources else ""
    return [
        p.title,
        p.org_name,
        ", ".join(p.region) if p.region else "",
        p.deadline.astimezone(KST).strftime("%Y-%m-%d") if p.deadline else "",
        "",                                   # 수업 일정 — 본문에서 따로 뽑지 않는다
        p.body_excerpt or p.one_line_summary or "",
        "",                                   # 수업 대상
        "",                                   # 모집 인원
        "\n".join(p.qualifications) if p.qualifications else "",
        "",                                   # 제출 서류
        url,
        "",                                   # 지원서 링크
        "",                                   # 지원 이메일 — 본문에서 개인정보를 지워 보관한다
    ]


def build_workbook(store: Store, now: datetime) -> tuple[bytes, int]:
    """(엑셀 바이트, 담은 줄 수). 템플릿의 '공고' 시트에 줄을 채운다."""
    wb = load_workbook(BytesIO(_template_bytes()))
    ws = wb[SHEET]

    header = [ws.cell(1, c).value for c in range(1, len(COLUMNS) + 1)]
    if header != COLUMNS:
        # 템플릿이 바뀌었는데 우리가 모르고 있는 상태다. 엉뚱한 칸에 값을 넣느니 멈춘다
        raise ValueError(f"양식의 열이 바뀌었다: {header}")

    rows = uploadable(store, now)
    for i, p in enumerate(rows, start=2):
        for c, value in enumerate(_row(p), start=1):
            ws.cell(i, c, value)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


def workbook_name(now: datetime) -> str:
    return f"강사잇다_공고업로드_{now.astimezone(KST):%Y-%m-%d}.xlsx"
