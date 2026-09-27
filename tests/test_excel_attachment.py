"""메일에 붙는 엑셀과 첨부 구조.

본문은 오늘 볼 것만 싣고 전체 목록은 첨부로 보낸다. 첨부가 깨지면 전체 목록을
볼 길이 없어지므로 시트 구성과 첨부 MIME 구조를 고정해 둔다.
"""
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from gia.config import ConfigBundle
from gia.extract.deadline import KST
from gia.models import Status
from gia.notify.email import SmtpConfig, build_message
from gia.report.build import select_postings
from gia.report.excel import build_workbook, workbook_name
from gia.store import Store
from tests.conftest import make_posting, make_source

NOW = datetime(2026, 9, 27, 9, 0, tzinfo=KST)


def _bundle(settings):
    src = make_source("board", "university", type="html_list",
                      list_url="https://site.example.org/board/list", keywords=["강사"])
    return ConfigBundle(settings=settings, sources=[src], aliases={})


def _filled(settings, tmp_path):
    store = Store(tmp_path)
    soon = make_posting("내일 마감 강사 모집", deadline=NOW + timedelta(days=1), now=NOW)
    later = make_posting("여유 있는 강사 모집", deadline=NOW + timedelta(days=40), now=NOW)
    for p in (soon, later):
        p.status = Status.new
        store.upsert(p)
    bundle = _bundle(settings)
    bundle.settings.report.closing_soon_days = 3
    return bundle, store


def test_workbook_has_four_sheets(settings, tmp_path):
    bundle, store = _filled(settings, tmp_path)
    data = select_postings(bundle, store, NOW)
    wb = load_workbook(BytesIO(build_workbook(data, store, bundle, NOW)))
    assert [n.split(" (")[0] for n in wb.sheetnames] == ["신규", "마감임박", "진행중 전체", "수집현황"]
    # 마감임박 시트에 내일 마감 건이 들어가야 한다
    ws = wb[[n for n in wb.sheetnames if n.startswith("마감임박")][0]]
    assert ws.cell(5, 5).value == "내일 마감 강사 모집"


def test_dday_is_a_formula_not_a_baked_number(settings, tmp_path):
    """받은 편지함에 며칠 묵혔다 열어도 그날 기준으로 세야 한다."""
    bundle, store = _filled(settings, tmp_path)
    data = select_postings(bundle, store, NOW)
    wb = load_workbook(BytesIO(build_workbook(data, store, bundle, NOW)))
    ws = wb[[n for n in wb.sheetnames if n.startswith("진행중")][0]]
    assert str(ws.cell(5, 8).value).startswith("=") and "TODAY()" in ws.cell(5, 8).value


def test_attachment_keeps_html_alternative(settings, tmp_path):
    """첨부를 붙여도 본문이 text/plain + text/html 로 남아야 한다."""
    bundle, store = _filled(settings, tmp_path)
    data = select_postings(bundle, store, NOW)
    blob = build_workbook(data, store, bundle, NOW)
    cfg = SmtpConfig(host="smtp.example.org", to=["a@example.org"], sender="b@example.org")
    msg = build_message(cfg, "제목", "<p>본문</p>", "본문", [(workbook_name(NOW), blob)])

    assert msg.get_content_maintype() == "multipart"
    types = {part.get_content_type() for part in msg.walk()}
    assert "text/plain" in types and "text/html" in types
    names = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    assert names == ["부울경_평생교육원_강사공고_2026-09-27.xlsx"]


def test_subject_carries_weekday_and_counts(settings, tmp_path):
    from gia.report.build import email_subject

    bundle, store = _filled(settings, tmp_path)
    data = select_postings(bundle, store, NOW)
    s = email_subject(data, NOW)
    assert "9/27(일)" in s and "신규 1건" in s and "마감임박 1건" in s
