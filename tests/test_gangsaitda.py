"""강사잇다 업로드 양식.

사이트가 받는 양식이라 열 이름이 하나라도 틀리면 파일 전체가 거부된다. 우리가
표를 흉내 내지 않고 받은 템플릿을 채우는 이유이고, 그 전제를 여기서 고정한다.
"""
from datetime import datetime, timedelta
from io import BytesIO

from openpyxl import load_workbook

from gia.extract.deadline import KST
from gia.models import Status
from gia.report.gangsaitda import COLUMNS, build_workbook, uploadable, workbook_name
from gia.store import Store
from tests.conftest import make_posting

NOW = datetime(2026, 9, 28, 9, 0, tzinfo=KST)


def _store(tmp_path):
    s = Store(tmp_path)
    live = make_posting("살아 있는 강사 모집", deadline=NOW + timedelta(days=10), now=NOW)
    past = make_posting("이미 지난 강사 모집", deadline=NOW - timedelta(days=1), now=NOW)
    noday = make_posting("마감일 없는 강사 모집", now=NOW)
    for p in (live, past, noday):
        p.status = Status.new
        s.upsert(p)
    return s


def test_only_rows_that_can_be_uploaded(tmp_path):
    """양식 안내: 지난 날짜는 올라가지 않는다. 마감일은 필수 칸이라 빈 줄도 거부된다."""
    titles = [p.title for p in uploadable(_store(tmp_path), NOW)]
    assert titles == ["살아 있는 강사 모집"]


def test_header_and_sheets_are_untouched(tmp_path):
    wb = load_workbook(BytesIO(build_workbook(_store(tmp_path), NOW)[0]))
    assert wb.sheetnames == ["공고", "예시", "안내"]
    ws = wb["공고"]
    assert [ws.cell(1, c).value for c in range(1, len(COLUMNS) + 1)] == COLUMNS


def test_row_lands_in_the_right_columns(tmp_path):
    blob, n = build_workbook(_store(tmp_path), NOW)
    assert n == 1
    ws = load_workbook(BytesIO(blob))["공고"]
    row = {ws.cell(1, c).value: ws.cell(2, c).value for c in range(1, len(COLUMNS) + 1)}
    assert row["제목"] == "살아 있는 강사 모집"
    assert row["마감일"] == "2026-10-08"      # 안내: 2026-10-15 처럼 적는다
    assert row["원문 링크"]


def test_body_excerpt_fills_the_required_detail_column(tmp_path):
    """'상세 내용' 은 파란 칸(필수)이다. 비우면 그 줄이 거부된다."""
    s = Store(tmp_path)
    p = make_posting("본문 있는 강사 모집", deadline=NOW + timedelta(days=5), now=NOW)
    p.status = Status.new
    p.body_excerpt = "주 1회 3시간, 12주 과정입니다."
    s.upsert(p)
    ws = load_workbook(BytesIO(build_workbook(s, NOW)[0]))["공고"]
    col = {ws.cell(1, c).value: c for c in range(1, len(COLUMNS) + 1)}
    assert ws.cell(2, col["상세 내용"]).value == "주 1회 3시간, 12주 과정입니다."


def test_rejects_a_template_whose_columns_moved(tmp_path, monkeypatch):
    """템플릿이 바뀐 걸 모르고 엉뚱한 칸에 값을 넣느니 멈춘다."""
    import gia.report.gangsaitda as g

    wb = load_workbook(BytesIO(g._template_bytes()))
    wb["공고"].cell(1, 1, "제목(변경됨)")
    buf = BytesIO(); wb.save(buf)
    monkeypatch.setattr(g, "_template_bytes", lambda: buf.getvalue())
    try:
        g.build_workbook(_store(tmp_path), NOW)
    except ValueError as e:
        assert "열이 바뀌었다" in str(e)
    else:
        raise AssertionError("바뀐 양식을 그냥 채웠다")


def test_file_name_carries_the_date():
    assert workbook_name(NOW) == "강사잇다_공고업로드_2026-09-28.xlsx"
